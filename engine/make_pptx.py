# -*- coding: utf-8 -*-
"""
Office AI 助手 · PPT 生成引擎
--------------------------------------------------------------------------
把一份大纲生成 .pptx。完全本地生成，不联网。

用法：
    python make_pptx.py -i 大纲.json  -o 输出.pptx
    python make_pptx.py -i 大纲.txt   -o 输出.pptx   [--theme tech-cyan]
    python make_pptx.py --json "{\"title\":...}"     -o 输出.pptx

大纲文本格式（以 @ 开头的行是元信息，其余按标题/要点组织）：
    @title 2026 年 AI 办公趋势
    @subtitle 洞察与展望
    @theme  tech-cyan

    # 背景与动机
    - 第一条要点
    - 第二条要点

    # 现状分析
    - ...

大纲 JSON 格式：
    {
      "title": "...", "subtitle": "...", "theme": "tech-cyan",
      "slides": [ {"title": "...", "bullets": ["...", "..."], "note": "备注"} ]
    }
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

APP_TITLE = "Office AI 助手"

# 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

FONT_CN = "Microsoft YaHei UI"
FONT_EN = "Segoe UI"

THEMES = {
    "tech-cyan": {
        "dark": "0F172A", "panel": "152036", "accent": "22D3EE",
        "title": "FFFFFF", "sub": "94A3B8", "text": "1F2937",
        "muted": "6B7280", "light": "FFFFFF", "line": "E5E7EB",
    },
    "business-blue": {
        "dark": "0B2447", "panel": "12305A", "accent": "4D8BF5",
        "title": "FFFFFF", "sub": "A8C0E0", "text": "1F2937",
        "muted": "6B7280", "light": "FFFFFF", "line": "E5E7EB",
    },
    "warm-amber": {
        "dark": "2B1B10", "panel": "3B2717", "accent": "F0A62B",
        "title": "FFFFFF", "sub": "D8BC9C", "text": "2B1B10",
        "muted": "7A6A5A", "light": "FFFDF9", "line": "EFE3D3",
    },
    "forest-green": {
        "dark": "0F2A22", "panel": "16382E", "accent": "34D399",
        "title": "FFFFFF", "sub": "9CC9B8", "text": "16302A",
        "muted": "6B7280", "light": "FFFFFF", "line": "E3EFE9",
    },
}

DEFAULT_THEME = "tech-cyan"


# ---------------------------------------------------------------- 文本工具
def set_font(run, size=None, bold=None, color=None, cn=FONT_CN, en=FONT_EN):
    """python-pptx 的 font.name 只设 latin，中文要另外写 a:ea，否则中文会回退到默认字体。"""
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = RGBColor.from_string(color)
    f.name = en
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", cn)


def add_rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.25):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    set_font(r, size=size, bold=bold, color=color)
    return tb


# ---------------------------------------------------------------- 版面
def slide_title(prs, t, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, spec["dark"])
    # 左侧强调条
    add_rect(s, Inches(0), Inches(2.9), Inches(0.16), Inches(1.5), spec["accent"])
    add_text(s, Inches(0.95), Inches(2.85), Inches(11.4), Inches(1.3),
             t["title"], 40, spec["title"], bold=True)
    if t.get("subtitle"):
        add_text(s, Inches(0.98), Inches(4.25), Inches(11.4), Inches(0.7),
                 t["subtitle"], 17, spec["sub"])
    add_text(s, Inches(0.98), Inches(6.55), Inches(11.4), Inches(0.4),
             t.get("footer") or "", 11, spec["sub"])
    return s


def slide_section(prs, idx, title, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, spec["panel"])
    add_rect(s, Inches(0.9), Inches(3.05), Inches(0.16), Inches(1.1), spec["accent"])
    add_text(s, Inches(1.3), Inches(2.95), Inches(1.6), Inches(0.7),
             "%02d" % idx, 30, spec["accent"], bold=True)
    add_text(s, Inches(1.32), Inches(3.62), Inches(10.6), Inches(0.9),
             title, 26, spec["title"], bold=True)
    return s


def slide_content(prs, idx, page, slide, spec, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, spec["light"])

    # 顶部：强调条 + 标题
    add_rect(s, Inches(0.7), Inches(0.62), Inches(0.9), Inches(0.09), spec["accent"])
    add_text(s, Inches(0.7), Inches(0.85), Inches(11.9), Inches(0.8),
             slide.get("title") or ("第 %d 页" % page), 25, spec["text"], bold=True)

    # 分隔线
    ln = add_rect(s, Inches(0.7), Inches(1.72), Inches(11.93), Emu(9525), spec["line"])

    bullets = slide.get("bullets") or []
    y = Inches(2.05)
    for b in bullets:
        level = 0
        text = b
        m = re.match(r"^(\s+)(.*)$", b)
        if m:
            level = min(len(m.group(1)) // 2, 2)
            text = m.group(2).lstrip("•·-* ").strip()
        else:
            text = b.lstrip("•·-* ").strip()
        if not text:
            continue
        indent = Inches(0.7 + level * 0.45)
        dot = add_rect(s, indent, y + Inches(0.115), Inches(0.075), Inches(0.075),
                       spec["accent"], shape=MSO_SHAPE.OVAL)
        add_text(s, indent + Inches(0.22), y, Inches(11.6 - level * 0.45), Inches(0.5),
                 text, 15 if level == 0 else 13.5,
                 spec["text"] if level == 0 else spec["muted"],
                 line_spacing=1.3)
        y += Inches(0.62 if level == 0 else 0.5)
        if y > Inches(6.6):
            break

    # 页脚
    add_text(s, Inches(0.7), Inches(6.92), Inches(6), Inches(0.3),
             slide.get("footer") or "", 10, spec["muted"])
    add_text(s, Inches(11.2), Inches(6.92), Inches(1.43), Inches(0.3),
             "%d / %d" % (page, total), 10, spec["muted"], align=PP_ALIGN.RIGHT)

    note = slide.get("note")
    if note:
        s.notes_slide.notes_text_frame.text = note
    return s


# ---------------------------------------------------------------- 解析
def parse_outline_text(text):
    deck = {"title": "", "subtitle": "", "theme": DEFAULT_THEME, "slides": []}
    cur = None
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.lstrip().startswith("@"):
            m = re.match(r"^@(\w+)\s+(.*)$", line.strip())
            if m:
                k, v = m.group(1).lower(), m.group(2).strip()
                if k in ("title", "subtitle", "theme", "footer"):
                    deck[k] = v
            continue
        if line.lstrip().startswith("#"):
            cur = {"title": line.lstrip("#").strip(), "bullets": [], "note": ""}
            deck["slides"].append(cur)
            continue
        if line.lstrip().startswith((">", "备注:", "note:")):
            if cur is not None:
                cur["note"] = line.lstrip("> ").replace("备注:", "").replace("note:", "").strip()
            continue
        if cur is None:
            cur = {"title": "", "bullets": [], "note": ""}
            deck["slides"].append(cur)
        # 保留前导空格（用于区分层级），只去掉行尾空白
        cur["bullets"].append(raw.rstrip())
    return deck


def load_deck(args):
    if args.json:
        return json.loads(args.json)
    path = args.input
    with open(path, encoding="utf-8-sig") as fh:
        raw = fh.read()
    if path.lower().endswith(".json"):
        return json.loads(raw)
    return parse_outline_text(raw)


# ---------------------------------------------------------------- 入口
def build(deck, out_path, theme_override=None):
    theme = theme_override or deck.get("theme") or DEFAULT_THEME
    spec = THEMES.get(theme) or THEMES[DEFAULT_THEME]

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    title = deck.get("title") or "演示文稿"
    slides = deck.get("slides") or []

    # 没有显式标题页信息时，用第一个幻灯片标题兜底
    if not deck.get("title") and slides:
        title = slides[0].get("title") or title

    slide_title(prs, {"title": title, "subtitle": deck.get("subtitle", ""),
                      "footer": deck.get("footer", "")}, spec)

    total = len(slides)
    for i, sl in enumerate(slides, 1):
        kind = (sl.get("type") or "").lower()
        if kind == "section":
            slide_section(prs, i, sl.get("title") or "", spec)
        else:
            slide_content(prs, i, i, sl, spec, total)

    prs.save(out_path)
    return {"out": out_path, "slides": total + 1, "theme": theme,
            "size": os.path.getsize(out_path)}


def main():
    ap = argparse.ArgumentParser(prog="make_pptx")
    ap.add_argument("-i", "--input", default=None)
    ap.add_argument("--json", default=None)
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--theme", default=None)
    args = ap.parse_args()

    if not args.input and not args.json:
        print(json.dumps({"ok": False, "error": "需要 -i 或 --json"}, ensure_ascii=False))
        return 2
    try:
        deck = load_deck(args)
        res = build(deck, args.output, args.theme)
        print(json.dumps({"ok": True, "result": res}, ensure_ascii=False))
        return 0
    except Exception as e:  # noqa: BLE001
        import traceback
        print(json.dumps({"ok": False, "error": str(e),
                          "trace": traceback.format_exc()}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    sys.exit(main())
